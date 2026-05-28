from sqlalchemy.orm import Session
from fastapi import HTTPException, status
from openpyxl import Workbook
from io import BytesIO
from fastapi.responses import StreamingResponse
from ..repositories import vehiculo_repository, dependencia_repository
from ..schemas.vehiculo_schema import VehiculoCreate, VehiculoUpdate

def get_all(db: Session):
    return vehiculo_repository.get_all(db)

def get_paginated(db, page, page_size, filter_field, filter_value, sort_field, sort_order, registro_marca, dependencia_id,
    situacion_id, fecha_desde, fecha_hasta):
    return vehiculo_repository.get_paginated_vehiculos(
        db=db,
        page=page,
        page_size=page_size,
        filter_field=filter_field,       
        filter_value=filter_value,
        sort_field=sort_field,
        sort_order=sort_order,
        registro_marca=registro_marca,
        dependencia_id=dependencia_id,
        situacion_id=situacion_id,
        fecha_desde=fecha_desde,
        fecha_hasta=fecha_hasta
    )

def get_by_id(db: Session, id: int):
    item = vehiculo_repository.get_by_id(db, id)
    if not item:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail=f"Vehículo con id={id} no encontrado"
        )
    return item

def create(db: Session, vehiculo: VehiculoCreate):
    # 1. Validación de Registro Nro (solo si se envió explícitamente)
    if vehiculo.registro_nro is not None:
        existingRO = vehiculo_repository.get_by_registroNro(db, vehiculo.registro_nro)
        if existingRO:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail=f"El vehículo con RO: '{vehiculo.registro_nro}' ya existe"
            )

    # 2. Validación de Dominio
    existingDominio = vehiculo_repository.get_by_Dominio(db, vehiculo.dominio)
    if existingDominio:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail=f"El vehículo con Dominio: '{vehiculo.dominio}' ya existe"
        )
    
    # 3. Validación de NroChasis
    chasis_nro = (vehiculo.chasis_nro or "").strip()
    existingChasis = vehiculo_repository.get_by_NroChasis(db, chasis_nro)
    if existingChasis:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail=f"El vehículo con Chasis: '{vehiculo.chasis_nro}' ya existe"
        )
    
    # 4. Validación de Motor
    motor_nro = (vehiculo.motor_nro or "").strip()
    existingMotor = vehiculo_repository.get_by_NroMotor(db, motor_nro)
    if existingMotor:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail=f"El vehículo con Motor: '{vehiculo.motor_nro}' ya existe"
        )
    
    # 3. Dependencia (OPCIONAL)
    dependencia = None
    if vehiculo.dependencia_id is not None:
        dependencia = dependencia_repository.get_by_id(
            db, vehiculo.dependencia_id
        )
        if not dependencia:
            raise HTTPException(
                status_code=status.HTTP_404_NOT_FOUND,
                detail="Dependencia no encontrada"
            )

    return vehiculo_repository.create(db, vehiculo, dependencia)


def update(db: Session, id: int, vehiculo: VehiculoUpdate):    
    # 1. Validación de Dominio
    existingDominio = vehiculo_repository.get_by_Dominio(db, vehiculo.dominio, exclude_id = id)
    if existingDominio:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail=f"El vehículo con Dominio: '{vehiculo.dominio}' ya existe"
        )
    # 2. Validación de NroChasis
    chasis_nro = (vehiculo.chasis_nro or "").rstrip()
    if chasis_nro and chasis_nro.upper() != "SIN INFORMAR":
        existingChasis = vehiculo_repository.get_by_NroChasis(db, chasis_nro, exclude_id = id)
        if existingChasis:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail=f"El vehículo con Chasis: '{chasis_nro}' ya existe"
            )
    
    # 3. Validación de Motor
    motor_nro = (vehiculo.motor_nro or "").rstrip()
    if motor_nro and motor_nro.upper() != "SIN INFORMAR":
        existingMotor = vehiculo_repository.get_by_NroMotor(db, motor_nro,  exclude_id = id)
        if existingMotor:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail=f"El vehículo con Motor: '{motor_nro}' ya existe"
            )
    
    db_item = vehiculo_repository.update(db, id, vehiculo)
    if not db_item:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail=f"El vehículo con id={id} no encontrado"
        )
    return db_item

def delete(db: Session, id: int):
    deleted = vehiculo_repository.delete(db, id)
    if not deleted:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail=f"El vehículo con id={id} no encontrado"
        )
    return {"message": f"El vehículo con id={id} eliminado correctamente"}

# Consultas

def get_by_dependencia(db: Session):
    return vehiculo_repository.get_by_dependencia(db)

def get_by_anio(db: Session):
    return vehiculo_repository.get_by_anio(db)

def get_by_condicion(db: Session):
    return vehiculo_repository.get_by_condicion(db)

def get_by_situacion(db: Session):
    return vehiculo_repository.get_by_situacion(db)

def get_historial_dependencias(db: Session, vehiculo_id: int):
    return vehiculo_repository.get_historial_dependencias(db, vehiculo_id)

def report_vehiculo_dependencia_activos(db: Session):

    wb = Workbook(write_only=True)
    ws = wb.create_sheet("Vehículos")

    ws.append(["Dominio", "Modelo", "Año", "Situacion", "Dependencia"])

    rows = vehiculo_repository.get_Vehiculos_dependencia_activos(db)

    for registro_nro, dominio, modelo_descripcion, anio, situacion_descripcion, dependencia_id, dependencia in rows:
        ws.append([
            dominio,
            modelo_descripcion,
            anio,
            situacion_descripcion,
            dependencia
        ])

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={
            "Content-Disposition":
            "attachment; filename=vehiculos_dependencia.xlsx"
        }
    )

def get_report_vehiculo_excel(db, filter_field, filter_value, sort_field, sort_order, registro_marca, dependencia_id,
    situacion_id, fecha_desde, fecha_hasta):

    wb = Workbook(write_only=True)
    ws = wb.create_sheet("Vehículos")

    ws.append(["Dominio", "Situacion", "Marca", "Modelo", "Año", "Dependencia"])

    rows = vehiculo_repository.get_paginated_vehiculos_report(db, filter_field, filter_value, sort_field, sort_order, registro_marca, dependencia_id,
    situacion_id, fecha_desde, fecha_hasta)

    for v in rows:
        ws.append([
            v.dominio,
            v.tipo_situacion.descripcion if v.tipo_situacion else None,
            v.tipo_marca.descripcion if v.tipo_marca else None,
            v.tipo_modelo.descripcion if v.tipo_modelo else None,
            v.anio if v.anio else None,
            v.dependencia.dependencia_descripcion if v.dependencia else None,
        ])

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={
            "Content-Disposition":
            "attachment; filename=vehiculos_reporte.xlsx"
        }
    )